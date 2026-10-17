import streamlit as st
import google.generativeai as genai
import os
import re
import io
import json
import urllib.parse
from youtube_search import YoutubeSearch
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Configuração da página e identidade visual CazéTV
st.set_page_config(
    page_title="CazéTV — Decupador de VT",
    page_icon="⚡",
    layout="wide"
)

st.markdown("""
<style>
    .stApp { background-color: #0b0f19; color: #f1f5f9; }
    .main-header { font-size: 2.2rem; font-weight: 800; color: #e11d48; margin-bottom: 0.2rem; text-transform: uppercase; }
    .sub-header { color: #94a3b8; margin-bottom: 1.5rem; font-size: 0.95rem; }
    .stTextArea textarea { background-color: #1e293b; color: #ffffff; border: 1px solid #334155; border-radius: 8px; font-size: 14px; line-height: 1.5; }
    .stButton button { background-color: #e11d48; color: white; font-weight: 700; font-size: 1rem; border-radius: 8px; padding: 0.7rem 1.5rem; border: none; width: 100%; transition: all 0.2s; }
    .stButton button:hover { background-color: #be123c; color: white; transform: translateY(-1px); }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-header">⚡ CAZÉTV — DECUPADOR DE VT</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-header">Autonomia editorial de IA: curadoria oficial de vídeos com timecode e pesquisa de fotos de imprensa.</div>', unsafe_allow_html=True)

# Chave de API
api_key = st.secrets.get("GEMINI_API_KEY", "")
if not api_key:
    with st.sidebar:
        st.subheader("⚙️ Configurações")
        api_key = st.text_input("Gemini API Key:", type="password")

# Lê a base de conhecimento real
base_conhecimento = ""
if os.path.exists("exemplos_cazetv.txt"):
    with open("exemplos_cazetv.txt", "r", encoding="utf-8", errors="ignore") as f:
        base_conhecimento = f.read()

# Insere hiperlink clicável no Word / Google Docs
def add_hyperlink(paragraph, url, text, color="E11D48"):
    try:
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(
            f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" r:id="{r_id}" w:history="1">'
            f'<w:r>'
            f'<w:rPr><w:color w:val="{color}"/><w:u w:val="single"/><w:b/></w:rPr>'
            f'<w:t>{text}</w:t>'
            f'</w:r>'
            f'</w:hyperlink>'
        )
        paragraph._p.append(hyperlink)
    except Exception:
        r = paragraph.add_run(f"{text} ({url})")
        r.font.color.rgb = RGBColor(225, 29, 72)
        r.font.underline = True

# Busca de vídeo oficial no YouTube
def buscar_video_youtube(termo_busca, timecode_str):
    try:
        segundos = 0
        match_tc = re.search(r'(\d{1,2}):(\d{2})', timecode_str)
        if match_tc:
            minutos = int(match_tc.group(1))
            segs = int(match_tc.group(2))
            segundos = (minutos * 60) + segs

        results = YoutubeSearch(termo_busca, max_results=3).to_dict()
        if results:
            video = results[0]
            url = f"https://www.youtube.com{video['url_suffix']}"
            if segundos > 0:
                url += f"&t={segundos}s"
            
            titulo = video.get('title', termo_busca)
            if len(titulo) > 65:
                titulo = titulo[:62] + "..."
            canal = video.get("channel", "YouTube")
            
            return {
                "tipo": "video",
                "titulo": f"▶️ {titulo} ({canal})",
                "url": url,
                "texto_markdown": f"[▶️ {titulo} ({canal})]({url})"
            }
    except Exception:
        pass
    return {
        "tipo": "video",
        "titulo": f"▶️ {termo_busca}",
        "url": "",
        "texto_markdown": f"**{termo_busca}**"
    }

# Busca de foto jornalística / Google Imagens
def gerar_link_foto(termo_busca, arquivo_sugerido):
    query_encoded = urllib.parse.quote(termo_busca)
    url_foto = f"https://www.google.com/search?tbm=isch&q={query_encoded}"
    nome = arquivo_sugerido if arquivo_sugerido else termo_busca
    return {
        "tipo": "foto",
        "titulo": f"📷 FOTO: {nome}",
        "url": url_foto,
        "texto_markdown": f"[📷 FOTO: {nome}]({url_foto})"
    }

# Extrator blindado que nunca quebra
def extrair_dados_resposta(texto_bruto):
    # 1. Tenta extrair JSON delimitado
    match = re.search(r'(\{[\s\S]*\})', texto_bruto)
    if match:
        try:
            dados = json.loads(match.group(1).strip())
            if "linhas" in dados and isinstance(dados["linhas"], list):
                return dados["linhas"]
        except Exception:
            pass
            
    # 2. Fallback: Se a IA devolveu tabela Markdown direta
    linhas_tabela = []
    for linha in texto_bruto.split("\n"):
        linha = linha.strip()
        if linha.startswith("|") and not re.search(r'\|\s*:?---', linha):
            colunas = [c.strip() for c in linha.split("|")[1:-1]]
            if len(colunas) >= 3:
                if "PROPOSTA" in colunas[0].upper():
                    continue
                tipo = "foto" if "FOTO" in colunas[2].upper() else "video"
                linhas_tabela.append({
                    "proposta": colunas[0],
                    "off": colunas[1],
                    "tipo": tipo,
                    "termo_busca": colunas[2],
                    "arquivo_sugerido": "FOTO_ARQUIVO.jpg",
                    "timecode": "",
                    "observacoes": ""
                })
    return linhas_tabela

# Gerador DOCX oficial
def gerar_documento_docx(off_completo, itens_decupados):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    h1 = doc.add_paragraph()
    r1 = h1.add_run("CAZÉTV — ROTEIRO E DECUPAGEM DE VT")
    r1.font.name = "Arial"
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(225, 29, 72)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    h2 = doc.add_paragraph()
    r2 = h2.add_run("1. LOCUÇÃO COMPLETA (OFF)")
    r2.font.name = "Arial"
    r2.font.size = Pt(13)
    r2.font.bold = True

    p_off = doc.add_paragraph()
    r_off = p_off.add_run(off_completo)
    r_off.font.name = "Georgia"
    r_off.font.size = Pt(11)
    r_off.font.italic = True
    p_off.paragraph_format.line_spacing = 1.3

    doc.add_paragraph()

    h3 = doc.add_paragraph()
    r3 = h3.add_run("2. TABELA DE DECUPAGEM TÉCNICA")
    r3.font.name = "Arial"
    r3.font.size = Pt(13)
    r3.font.bold = True

    tabela = doc.add_table(rows=len(itens_decupados) + 1, cols=3)
    tabela.autofit = False
    larguras = [Inches(3.0), Inches(1.8), Inches(2.7)]

    headers = ["PROPOSTA (IMAGEM / VÍDEO / ÁUDIO)", "OFF", "LINKS / OBSERVAÇÕES E TIMECODES"]
    for j, h_text in enumerate(headers):
        cell = tabela.rows[0].cells[j]
        cell.width = larguras[j]
        p = cell.paragraphs[0]
        p.text = h_text
        shading = parse_xml(r'<w:shd {} w:fill="111827"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        for r in p.runs:
            r.font.name = "Arial"
            r.font.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(255, 255, 255)

    for i, item in enumerate(itens_decupados):
        row = tabela.rows[i + 1]
        
        # Coluna 1
        c0 = row.cells[0]
        c0.width = larguras[0]
        p0 = c0.paragraphs[0]
        p0.text = item.get("proposta", "")
        for r in p0.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)

        # Coluna 2
        c1 = row.cells[1]
        c1.width = larguras[1]
        p1 = c1.paragraphs[0]
        p1.text = item.get("off", "")
        for r in p1.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)

        # Coluna 3 com Hiperlink Nativo
        c2 = row.cells[2]
        c2.width = larguras[2]
        p2 = c2.paragraphs[0]
        
        info_link = item.get("link_info", {})
        if info_link.get("url"):
            cor = "2563EB" if item.get("tipo") == "foto" else "E11D48"
            add_hyperlink(p2, info_link["url"], info_link["titulo"], color=cor)
            
            detalhes = ""
            if item.get("tipo") == "video" and item.get("timecode"):
                detalhes += f"\n({item.get('timecode')})"
            if item.get("observacoes"):
                detalhes += f" - {item.get('observacoes')}"
            
            r_obs = p2.add_run(detalhes)
            r_obs.font.name = "Arial"
            r_obs.font.size = Pt(9)
        else:
            p2.text = f"{item.get('termo_busca', '')} - {item.get('observacoes', '')}"

    for row in tabela.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders {}><w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
                r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
                r'<w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>'.format(nsdecls('w'))
            )
            tcPr.append(tcBorders)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

off_text = st.text_area("LOCUÇÃO / TEXTO DO OFF:", height=220, placeholder="Cole o texto corrido do OFF aqui...")

if off_text:
    palavras = len(off_text.split())
    segundos = int((palavras / 135) * 60)
    st.caption(f"📊 **{palavras} palavras** | Estimativa de locução: **~{segundos} segundos**")

if st.button("🎬 GERAR DECUPAGEM COM VÍDEOS E FOTOS"):
    if not api_key:
        st.error("⚠️ Chave de API não configurada.")
    elif not off_text.strip():
        st.warning("⚠️ Insira o texto do OFF antes de gerar a decupagem.")
    else:
        with st.spinner("Analisando proposta, decidindo entre vídeos e fotos e pesquisando fontes oficiais..."):
            try:
                genai.configure(api_key=api_key)
                
                modelos_validos = []
                try:
                    for m in genai.list_models():
                        if 'generateContent' in m.supported_generation_methods:
                            modelos_validos.append(m.name)
                except Exception:
                    pass

                if not modelos_validos:
                    modelos_validos = ["models/gemini-2.5-flash", "models/gemini-3.8-flash", "gemini-2.5-flash"]

                system_prompt = f"""
                VOCÊ É O DIRETOR CHEFE DE VTS E DECUPADOR DA CAZÉTV.
                Sua função é gerar a decupagem oficial da CazéTV a partir do OFF.

                Base de Conhecimento CazéTV:
                {base_conhecimento}

                SUA RESPOSTA DEVE SER ESTRITAMENTE UM JSON VÁLIDO no seguinte formato (sem introduções ou conversas):
                {{
                  "linhas": [
                    {{
                      "proposta": "Descrição visual detalhada com planos, slows, letterings estilizados e sobe som de torcida/áudio",
                      "off": "Trecho curto exato correspondente da locução",
                      "tipo": "video" ou "foto",
                      "termo_busca": "Termos de busca cirúrgicos incluindo a fonte oficial ou contexto",
                      "arquivo_sugerido": "Nome do arquivo (ex: FOTO_FILIPE_LUIS_MONACO.jpg)",
                      "timecode": "MM:SS - MM:SS (se for vídeo)",
                      "observacoes": "Instruções de edição, enquadramento ou áudio"
                    }}
                  ]
                }}

                DIRETRIZES DE DECISÃO DA IA:
                1. Você tem total autonomia para decidir se cada momento pede VÍDEO ou FOTO:
                   - Se pedir VÍDEO: inclua no termo_busca o canal oficial mais relevante (ex: canal do clube, liga, emissora detentora) e forneça a janela de timecode (MM:SS - MM:SS).
                   - Se pedir FOTO: gere uma busca cirúrgica para portais de imprensa/Google Imagens e defina o nome do arquivo interno (arquivo_sugerido).
                """

                prompt_final = f"Decupe este OFF em JSON:\n\n{off_text}"

                response = None
                for modelo in modelos_validos:
                    try:
                        model = genai.GenerativeModel(
                            model_name=modelo,
                            generation_config={"temperature": 0.2},
                            system_instruction=system_prompt
                        )
                        response = model.generate_content(prompt_final)
                        if response and response.text:
                            break
                    except Exception:
                        continue

                if response and response.text:
                    itens_decupados = extrair_dados_resposta(response.text)
                    
                    if not itens_decupados:
                        st.error("A IA não retornou linhas válidas. Tente novamente.")
                    else:
                        tabela_markdown = "| PROPOSTA (IMAGEM / VÍDEO / ÁUDIO) | OFF | LINKS / OBSERVAÇÕES E TIMECODES |\n| :--- | :--- | :--- |\n"
                        
                        for item in itens_decupados:
                            tipo = item.get("tipo", "video")
                            
                            if tipo == "foto":
                                termo = item.get("termo_busca") or item.get("arquivo_sugerido") or "foto futebol"
                                info_link = gerar_link_foto(termo, item.get("arquivo_sugerido", termo))
                                item["link_info"] = info_link
                                obs = item.get("observacoes", "")
                                coluna3 = f"{info_link['texto_markdown']}<br>*(Observação: {obs})*"
                            else:
                                info_link = buscar_video_youtube(item.get("termo_busca", ""), item.get("timecode", ""))
                                item["link_info"] = info_link
                                tc = item.get("timecode", "")
                                obs = item.get("observacoes", "")
                                coluna3 = f"{info_link['texto_markdown']}<br>({tc}) - {obs}"
                            
                            prop = item.get("proposta", "").replace("\n", "<br>")
                            off_trecho = item.get("off", "").replace("\n", " ")
                            tabela_markdown += f"| {prop} | {off_trecho} | {coluna3} |\n"

                        st.success("✅ Decupagem completa gerada com vídeos oficiais e fotos!")
                        st.markdown("---")
                        st.markdown(tabela_markdown, unsafe_allow_html=True)

                        docx_buffer = gerar_documento_docx(off_text, itens_decupados)

                        st.download_button(
                            label="📄 BAIXAR EM FORMATO GOOGLE DOCS / WORD (.DOCX)",
                            data=docx_buffer,
                            file_name="decupagem_cazetv.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                else:
                    st.error("Não foi possível processar a resposta da IA.")
                    
            except Exception as e:
                st.error(f"Erro geral: {e}")
